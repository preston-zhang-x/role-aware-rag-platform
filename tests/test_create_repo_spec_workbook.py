from pathlib import Path

from openpyxl import load_workbook

from script import create_repo_spec_workbook as workbook_script


def workbook_text_values(workbook) -> str:
    values: list[str] = []
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    values.append(cell.value)
    return "\n".join(values)


def test_collect_repo_spec_data_aligns_rag_auth_and_roles():
    spec = workbook_script.collect_repo_spec_data()

    endpoint_map = {endpoint.if_id: endpoint for endpoint in spec.endpoints}
    feature_map = {feature.feature_id: feature for feature in spec.features}
    screen_map = {screen.screen_id: screen for screen in spec.screens}
    item_map = {item.item_id: item for item in spec.screen_items}
    rule_map = {rule.rule_id: rule for rule in spec.rules}

    assert endpoint_map["IF-008"].auth_kind == "bearer"
    assert endpoint_map["IF-008"].auth_text == "Bearer 必須 / 認証済み"
    assert "current_user.role" in endpoint_map["IF-008"].note
    assert "current_user.role" in feature_map["FN-008"].note
    assert screen_map["SCR-003"].auth_text == "Bearer 必須"

    current_role_item = item_map["ITM-SCR003-002"]
    assert current_role_item.ui_type == "badge"
    assert current_role_item.editable_text == "読取専用"
    assert "current_user.role" in current_role_item.note
    assert "現在ログイン中ユーザーの role を表示" == current_role_item.rule_text

    assert "admin/manager" in rule_map["RL-003"].when_text
    assert "staff は不可" == rule_map["RL-003"].note
    assert "current_user.role" in rule_map["RL-009"].then_text

    joined = "\n".join(
        [
            *[feature.note for feature in spec.features],
            *[screen.summary for screen in spec.screens],
            *[item.note for item in spec.screen_items],
            *[rule.note for rule in spec.rules],
            *[rule.condition for rule in spec.rules],
            *[test_point.scenario for test_point in spec.test_points],
        ]
    )
    assert "body.role" not in joined
    assert "admin/editor/viewer" not in joined
    assert "viewer / editor / admin" not in joined


def test_write_repo_spec_workbook_generates_current_and_styled_workbook(tmp_path: Path):
    output_path = tmp_path / "repo_spec.xlsx"

    saved_path = workbook_script.write_repo_spec_workbook(output_path)

    assert saved_path == output_path
    workbook = load_workbook(saved_path)
    try:
        assert {"表紙・改定履歴", "画面設計_RAG問合せ", "認可マトリクス"} <= set(
            workbook.sheetnames
        )

        cover = workbook["表紙・改定履歴"]
        rag = workbook["画面設計_RAG問合せ"]
        auth_matrix = workbook["認可マトリクス"]

        assert cover["A9"].value == "KPI サマリー"
        assert cover["A1"].fill.fill_type == "solid"
        assert cover.sheet_view.showGridLines is False

        assert rag["C9"].value == "Current Role"
        assert rag["F9"].value == "JWT から解決した current_user.role"
        assert rag.sheet_view.showGridLines is False

        auth_headers = [
            auth_matrix.cell(row=5, column=col).value for col in range(1, 10)
        ]
        assert auth_headers == [
            "IF-ID",
            "名称",
            "Path",
            "匿名",
            "admin",
            "manager",
            "staff",
            "認可方式",
            "備考",
        ]

        if_008_row = None
        for row_index in range(6, auth_matrix.max_row + 1):
            if auth_matrix.cell(row=row_index, column=1).value == "IF-008":
                if_008_row = row_index
                break

        assert if_008_row is not None
        assert [
            auth_matrix.cell(row=if_008_row, column=col).value for col in range(4, 8)
        ] == ["N", "Y", "Y", "Y"]
        assert auth_matrix["A5"].border.left.style == "medium"

        all_text = workbook_text_values(workbook)
        assert "body.role" not in all_text
        assert "admin/editor/viewer" not in all_text
        assert "viewer / editor / admin" not in all_text
    finally:
        workbook.close()
