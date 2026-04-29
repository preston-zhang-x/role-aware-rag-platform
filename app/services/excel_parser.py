"""
日本式の Excel 仕様書を Markdown に変換するパーサー。

セル結合や方眼紙レイアウトを含むシートを対象に、見出し、表、キー・値形式の
情報をできるだけ読み取りやすい形へ整形する。
"""

import re
from contextlib import ExitStack
from dataclasses import dataclass
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.services.base_parser import BaseParser, ParseError
from app.services.parse_result import (
    ChunkMeta,
    ContentType,
    MarkdownEscaper,
    ParseResult,
    ParserWarning,
)

CellLike = Cell | MergedCell


@dataclass(frozen=True)
class RowCell:
    """行内で意味を持つセルを表す。

    属性:
        text: セルの内容テキスト（テキスト/KV用にマークダウンエスケープ済み）
        start_col: セルの開始列インデックス（0ベース）
        end_col: セルの終了列インデックス（0ベース）
    """

    text: str
    start_col: int
    end_col: int

    @property
    def col_span(self) -> int:
        return self.end_col - self.start_col + 1


@dataclass(frozen=True)
class RowProfile:
    """1 行分のセル配置をまとめた情報。

    属性:
        row_index: シート内の行番号（0ベース）
        raw_row: 行の生データ（全列分の値）
        cells: 処理済みセルオブジェクトのリスト
    """

    row_index: int
    raw_row: tuple[str, ...]
    cells: tuple[RowCell, ...]


class JapaneseExcelParser(BaseParser):
    """
    日本式 Excel 仕様書向けのパーサー。

    方眼紙形式のシートでよく使われる結合セルや、表とキー・値が混在する
    レイアウトを Markdown に変換する。
    """

    # レイアウト判定に使うしきい値。
    MAX_HEADING_LENGTH = 48  # 見出しとして扱うテキストの最大文字数
    MIN_TABLE_ROWS = 2  # テーブルと判定するための最小行数
    MAX_FORM_ROW_CELLS = 10  # KV形式行として扱う最大セル数
    MIN_HEADING_LENGTH = 2
    REPEATED_WIDE_LABEL_MIN_SPAN = 4
    MIN_SHARED_TABLE_COLUMNS = 3
    TWO_ROW_TABLE_SHARED_DENSITY = 0.9
    MIN_STRUCTURED_MERGED_COLUMNS = 4
    LEADING_CONTEXT_MIN_CELLS = 5
    LEADING_CONTEXT_MIN_SPAN = 3
    MAX_KV_KEY_COL_SPAN = 2
    KV_LABEL_TRAILING_CHARS = ":："
    NON_HEADING_PREFIXES = ("->", "→", "=>")
    NON_HEADING_EXACT_TEXTS = {
        "with_payload=True",
        "with_vectors=False",
        "source_file",
        "source_key",
        "allowed_roles",
        "chunk_index",
    }
    NOISE_HEADING_PATTERN = re.compile(
        r"^[\s()\[\]{}0-9０-９.．-]+$"
    )  # ノイズ見出しパターン（進捗番号など）

    def __init__(self, density_threshold: float = 0.5):
        """
        パーサーを初期化する。

        引数:
            density_threshold: 後方互換のために残している旧設定。
                現在のテーブル判定は列位置の揃い方を優先する。
        """
        self.density_threshold = density_threshold

    def can_handle(self, file_path: str) -> bool:
        """拡張子で判定: .xlsx, .xlsm"""
        suffix = Path(file_path).suffix.lower()
        return suffix in {".xlsx", ".xlsm"}

    def parse(self, file_path: str) -> ParseResult:
        """
        Excel ファイル全体を Markdown に変換する。

        引数:
            file_path: パース対象のExcelファイルのパス

        戻り値:
            ParseResult: 抽出されたマークダウンテキスト、チャンク、警告のセット

        例外:
            ParseError: ファイルが見つからない、または読み込みに失敗した場合
        """
        path = Path(file_path)
        if not path.exists():
            raise ParseError(f"ファイルが見つかりません: {file_path}", file_path)

        all_chunks: list[ChunkMeta] = []
        all_warnings: list[str] = []
        full_text = ""

        with ExitStack() as stack:
            # 表示値と数式を両方読む。計算済み値がないセルでは数式を残す。
            try:
                wb_cached = openpyxl.load_workbook(file_path, data_only=True)
                stack.callback(wb_cached.close)
                wb_formula = openpyxl.load_workbook(file_path, data_only=False)
                stack.callback(wb_formula.close)
            except Exception as exc:
                raise ParseError(f"Excelファイルの読み込みに失敗: {exc}", file_path)

            for sheet_name in wb_cached.sheetnames:
                ws_cached = wb_cached[sheet_name]
                ws_formula = wb_formula[sheet_name]

                # 非表示シートは検索対象に含めない。
                if ws_cached.sheet_state != "visible":
                    all_warnings.append(
                        f"{ParserWarning.HIDDEN_SHEET_SKIPPED.value}: {sheet_name}"
                    )
                    continue

                # シート名は後続チャンクの文脈になるため、見出しとして出力する。
                full_text, _ = self._append_section(full_text, f"# {sheet_name}")

                section_md, section_chunks, comments_md = self._parse_visible_sheet(
                    ws_cached,
                    ws_formula,
                    file_path,
                    sheet_name,
                    all_warnings,
                )
                if section_md:
                    full_text, section_start = self._append_section(
                        full_text, section_md
                    )
                    if section_start is not None:
                        for chunk in section_chunks:
                            chunk.char_start += section_start
                            chunk.char_end += section_start
                all_chunks.extend(section_chunks)

                # セルコメントは本文とは分けて末尾に追加する。
                if comments_md:
                    full_text, _ = self._append_section(full_text, comments_md)

        text = full_text + "\n" if full_text else ""
        return ParseResult(text=text, chunks=all_chunks, warnings=all_warnings)

    def _parse_visible_sheet(
        self,
        ws_cached: Worksheet,
        ws_formula: Worksheet,
        file_path: str,
        sheet_name: str,
        warnings: list[str],
    ) -> tuple[str, list[ChunkMeta], str]:
        """可視シートを本文 Markdown、チャンク、コメント Markdown へ変換する。"""
        grid, merge_spans = self._build_grid_with_merged_cells(
            ws_cached,
            ws_formula,
            warnings,
        )
        if not grid:
            return "_(空のシート)_", [], ""

        section_md, section_chunks = self._heuristic_scan(
            grid,
            file_path,
            sheet_name,
            merge_spans,
        )
        return section_md, section_chunks, self._extract_comments(ws_cached)

    def _append_section(
        self,
        full_text: str,
        section: str,
    ) -> tuple[str, int | None]:
        """最終 Markdown へセクションを追加し、その開始位置を返す。"""
        normalized = section.strip()
        if not normalized:
            return full_text, None

        separator = "\n\n" if full_text else ""
        start = len(full_text) + len(separator)
        return f"{full_text}{separator}{normalized}", start

    def _build_chunk_meta(
        self,
        file_path: str,
        content_type: ContentType,
        sheet_name: str | None = None,
        cell_range: str | None = None,
        is_broadcast_fill: bool = False,
    ) -> ChunkMeta:
        """チャンク生成時に共通して使うメタデータを組み立てる。"""
        return ChunkMeta(
            source_file=file_path,
            content_type=content_type,
            sheet_name=sheet_name,
            cell_range=cell_range,
            is_broadcast_fill=is_broadcast_fill,
        )

    def _build_grid_with_merged_cells(
        self,
        ws_cached: Worksheet,
        ws_formula: Worksheet,
        warnings: list[str],
    ) -> tuple[list[list[str]], dict[tuple[int, int], tuple[int, int]]]:
        """結合セルを考慮して、処理しやすい二次元グリッドを作る。

        戻り値:
            grid: 全セル値のフラット二次元リスト
            merge_spans: （アンカー座標 -> (行数, 列数)）のマッピング
        """
        if ws_cached.max_row is None or ws_cached.max_column is None:
            return [], {}

        max_row = ws_cached.max_row
        max_col = ws_cached.max_column
        grid: list[list[str]] = [["" for _ in range(max_col)] for _ in range(max_row)]

        # 表示値と数式を同じ範囲で読み、後段で結合セルの従属セルを空に戻す。
        cached_rows = ws_cached.iter_rows(
            min_row=1,
            max_row=max_row,
            min_col=1,
            max_col=max_col,
        )
        formula_rows = ws_formula.iter_rows(
            min_row=1,
            max_row=max_row,
            min_col=1,
            max_col=max_col,
        )
        for row_idx, (cached_row, formula_row) in enumerate(
            zip(cached_rows, formula_rows, strict=True)
        ):
            for col_idx, (cached_cell, formula_cell) in enumerate(
                zip(cached_row, formula_row, strict=True)
            ):
                grid[row_idx][col_idx] = self._read_cell_value(
                    cached_cell,
                    formula_cell,
                    warnings,
                )

        merge_spans: dict[tuple[int, int], tuple[int, int]] = {}
        for merged_range in ws_cached.merged_cells.ranges:
            # 結合セルの左上だけを値の持ち主として扱う。
            anchor = (merged_range.min_row - 1, merged_range.min_col - 1)
            merge_spans[anchor] = (
                merged_range.max_row - merged_range.min_row + 1,
                merged_range.max_col - merged_range.min_col + 1,
            )
            for row_idx in range(merged_range.min_row, merged_range.max_row + 1):
                for col_idx in range(merged_range.min_col, merged_range.max_col + 1):
                    if (row_idx - 1, col_idx - 1) != anchor:
                        grid[row_idx - 1][col_idx - 1] = ""

            warnings.append(
                f"{ParserWarning.MERGED_CELL_BROADCAST.value}: {merged_range}"
            )

        return grid, merge_spans

    def _read_cell_value(
        self,
        cached_cell: CellLike,
        formula_cell: CellLike,
        warnings: list[str],
    ) -> str:
        """
        セルの表示値を文字列として取得する。

        計算済みの値があればそれを使い、値がなく数式だけが残っている場合は
        数式テキストを保持する。

        引数:
            cached_cell: data_only=True で読み込んだセル（計算結果）
            formula_cell: data_only=False で読み込んだセル（数式）
            warnings: 警告メッセージを追記するリスト

        戻り値:
            セルの値を文字列化したもの
        """
        cached_value = cached_cell.value
        formula_value = formula_cell.value

        if cached_value is not None:
            return str(cached_value)

        if formula_value is not None and str(formula_value).startswith("="):
            warnings.append(f"{ParserWarning.FORMULA_NO_CACHE.value}: {formula_value}")
            return f"formula: {formula_value}"

        return ""

    def _heuristic_scan(
        self,
        grid: list[list[str]],
        file_path: str,
        sheet_name: str,
        merge_spans: dict[tuple[int, int], tuple[int, int]],
    ) -> tuple[str, list[ChunkMeta]]:
        """行の並びを走査し、テキスト、キー・値、表の単位に分ける。

        方眼紙 Excel では表とフォーム項目が同じシートに混ざるため、
        連続する行の列位置を見ながら Markdown の出力形式を選ぶ。

        戻り値:
            markdown: Markdown形式の文字列
            chunks: 抽出されたメタデータチャンク
        """
        effective_cols = self._calc_effective_cols(grid)
        if effective_cols == 0:
            return "", []

        row_profiles = [
            self._build_row_profile(row_idx, row[:effective_cols], merge_spans)
            for row_idx, row in enumerate(grid)
        ]

        markdown_parts: list[str] = []
        chunks: list[ChunkMeta] = []
        last_text_block: str | None = None
        index = 0

        while index < len(row_profiles):
            row = row_profiles[index]
            if not row.cells:
                index += 1
                continue

            if len(row.cells) == 1:
                block = self._render_text_row(
                    row, file_path, sheet_name, last_text_block
                )
                if block is not None:
                    markdown_parts.append(block[0])
                    chunks.append(block[1])
                    last_text_block = block[2]
                index += 1
                continue

            table_len = self._detect_table_length(row_profiles[index:])
            if table_len:
                table_rows = row_profiles[index : index + table_len]
                table_md, chunk = self._render_table_block(
                    table_rows,
                    file_path,
                    sheet_name,
                )
                if table_md:
                    markdown_parts.append(table_md)
                    chunks.append(chunk)
                last_text_block = None
                index += table_len
                continue

            form_md, chunk = self._render_form_row(
                row,
                row_profiles[index + 1] if index + 1 < len(row_profiles) else None,
                file_path,
                sheet_name,
            )
            if form_md:
                markdown_parts.append(form_md)
                chunks.append(chunk)
            last_text_block = None
            index += 1

        return self._with_chunk_positions(markdown_parts, chunks), chunks

    def _with_chunk_positions(
        self,
        markdown_parts: list[str],
        chunks: list[ChunkMeta],
    ) -> str:
        """シート内 Markdown を連結し、各チャンクの文字位置を更新する。"""
        section_text = ""
        for markdown, chunk in zip(markdown_parts, chunks, strict=False):
            if section_text:
                section_text += "\n\n"
            chunk.char_start = len(section_text)
            section_text += markdown
            chunk.char_end = len(section_text)
        return section_text

    def _build_row_profile(
        self,
        row_idx: int,
        row: list[str],
        merge_spans: dict[tuple[int, int], tuple[int, int]],
    ) -> RowProfile:
        cells: list[RowCell] = []
        for col_idx, raw_value in enumerate(row):
            value = MarkdownEscaper.escape_cell(raw_value)
            if not value:
                continue
            _, col_span = merge_spans.get((row_idx, col_idx), (1, 1))
            cell = RowCell(
                text=value,
                start_col=col_idx,
                end_col=col_idx + col_span - 1,
            )
            cells.append(cell)
        return RowProfile(row_index=row_idx, raw_row=tuple(row), cells=tuple(cells))

    def _detect_table_length(self, rows: list[RowProfile]) -> int:
        """先頭行から表として扱える範囲を検出する。

        行を 1 つずつ追加し、表としての形が崩れたところで止める。

        引数:
            rows: 対象となるRowProfileのリスト

        戻り値:
            連続するテーブル行の数（テーブルでない場合は0）
        """
        if len(rows) < self.MIN_TABLE_ROWS:
            return 0

        max_length = 0
        for end in range(self.MIN_TABLE_ROWS, len(rows) + 1):
            candidate = rows[:end]
            if self._looks_like_table_block(candidate):
                max_length = end
                continue
            if max_length:
                break
        return max_length

    def _looks_like_table_block(self, rows: list[RowProfile]) -> bool:
        """行グループが表として扱えるかを判定する。"""
        if len(rows) < self.MIN_TABLE_ROWS:
            return False
        if any(len(row.cells) < 2 for row in rows):
            return False

        first_cells = [row.cells[0] for row in rows[:2]]
        if (
            len(first_cells) == 2
            and first_cells[0].text == first_cells[1].text
            and first_cells[0].col_span >= self.REPEATED_WIDE_LABEL_MIN_SPAN
        ):
            return False

        if self._has_structured_merged_columns(rows):
            return True

        column_sets = [
            {cell.start_col for cell in row.cells} for row in rows[: min(3, len(rows))]
        ]
        shared_columns = set.intersection(*column_sets)
        average_width = sum(len(row.cells) for row in rows) / len(rows)

        if not shared_columns:
            return False

        shared_span = max(shared_columns) - min(shared_columns) + 1
        shared_density = len(shared_columns) / shared_span if shared_span else 0

        if len(rows) == 2:
            return (
                len(shared_columns) >= self.MIN_SHARED_TABLE_COLUMNS
                and shared_density >= self.TWO_ROW_TABLE_SHARED_DENSITY
            )

        return (
            len(shared_columns) >= self.MIN_SHARED_TABLE_COLUMNS
            and average_width >= self.MIN_SHARED_TABLE_COLUMNS
        )

    def _has_structured_merged_columns(self, rows: list[RowProfile]) -> bool:
        """結合幅のある列が、行をまたいで同じ位置に並んでいるかを見る。"""
        sample_rows = rows[: min(3, len(rows))]
        if len(sample_rows) < 2:
            return False

        if (
            min(len(row.cells) for row in sample_rows)
            < self.MIN_STRUCTURED_MERGED_COLUMNS
        ):
            return False

        if not any(cell.col_span > 1 for row in sample_rows for cell in row.cells):
            return False

        if self._looks_like_merged_kv_pairs(sample_rows):
            return False

        header_positions = tuple(cell.start_col for cell in sample_rows[0].cells)
        if len(header_positions) < self.MIN_STRUCTURED_MERGED_COLUMNS:
            return False

        return all(
            tuple(cell.start_col for cell in row.cells) == header_positions
            for row in sample_rows[1:]
        )

    def _looks_like_merged_kv_pairs(self, rows: list[RowProfile]) -> bool:
        """結合セルを使ったキー・値の並びかどうかを判定する。"""
        if not rows:
            return False

        cell_count = len(rows[0].cells)
        if (
            cell_count < 4
            or cell_count % 2 != 0
            or cell_count > self.MAX_FORM_ROW_CELLS
        ):
            return False

        reference_positions = tuple(cell.start_col for cell in rows[0].cells)
        if any(
            tuple(cell.start_col for cell in row.cells) != reference_positions
            for row in rows
        ):
            return False

        for row in rows:
            for index in range(0, len(row.cells), 2):
                key_cell = row.cells[index]
                value_cell = row.cells[index + 1]
                if key_cell.col_span > self.MAX_KV_KEY_COL_SPAN:
                    return False
                if value_cell.col_span <= key_cell.col_span:
                    return False

        return True

    def _render_text_row(
        self,
        row: RowProfile,
        file_path: str,
        sheet_name: str,
        last_text_block: str | None,
    ) -> tuple[str, ChunkMeta, str] | None:
        text = row.cells[0].text
        if self._should_skip_text_block(text, last_text_block):
            return None

        markdown = self._render_text_block(text)
        if not markdown:
            return None

        chunk = self._build_chunk_meta(
            file_path=file_path,
            content_type=ContentType.TEXT,
            sheet_name=sheet_name,
            cell_range=f"row {row.row_index + 1}",
            is_broadcast_fill=row.cells[0].col_span > 1,
        )
        return markdown, chunk, text

    def _should_skip_text_block(
        self,
        text: str,
        last_text_block: str | None,
    ) -> bool:
        """直前と同じ見出しや、番号だけの行を出力対象から外す。"""
        if text == last_text_block:
            return True
        if self.NOISE_HEADING_PATTERN.fullmatch(text):
            return True
        return False

    def _looks_like_heading(self, text: str) -> bool:
        """
        見出しとして扱える短いラベルかどうかを判定する。

        長文や注記は段落へ降格する。
        """
        normalized = text.strip()
        if (
            len(normalized) < self.MIN_HEADING_LENGTH
            or len(normalized) > self.MAX_HEADING_LENGTH
        ):
            return False
        if self.NOISE_HEADING_PATTERN.fullmatch(normalized):
            return False
        if normalized in self.NON_HEADING_EXACT_TEXTS:
            return False
        if normalized.startswith(self.NON_HEADING_PREFIXES):
            return False
        if any(token in normalized for token in ("。", "<br>", r"\|", ":", "：")):
            return False
        return True

    def _render_text_block(self, text: str) -> str:
        """単独セルの行を見出しまたは段落として整形する。"""
        if self._looks_like_heading(text):
            return f"## {text}"
        return text

    def _render_form_row(
        self,
        row: RowProfile,
        next_row: RowProfile | None,
        file_path: str,
        sheet_name: str,
    ) -> tuple[str, ChunkMeta]:
        """1 行のキー・値形式データを箇条書きに変換する。"""
        values = [cell.text for cell in row.cells]
        if self._should_drop_leading_context(row, next_row):
            values = values[1:]

        parts: list[str] = []
        i = 0
        while i < len(values):
            if i + 1 < len(values):
                label = self._clean_kv_label(values[i])
                parts.append(f"- **{label}:** {values[i + 1]}")
                i += 2
                continue
            parts.append(f"- {values[i]}")
            i += 1

        chunk = self._build_chunk_meta(
            file_path=file_path,
            content_type=ContentType.KEY_VALUE,
            sheet_name=sheet_name,
            cell_range=f"row {row.row_index + 1}",
        )
        return "\n".join(parts), chunk

    def _clean_kv_label(self, label: str) -> str:
        """キー・値形式のラベルから末尾の区切り記号を取り除く。"""
        cleaned = label.strip().rstrip(self.KV_LABEL_TRAILING_CHARS).strip()
        return cleaned or label

    def _should_drop_leading_context(
        self,
        row: RowProfile,
        next_row: RowProfile | None,
    ) -> bool:
        """広い先頭セルが行全体の見出しに見える場合は値ペアから外す。"""
        if len(row.cells) < self.LEADING_CONTEXT_MIN_CELLS or len(row.cells) % 2 == 0:
            return False

        first = row.cells[0]
        if first.col_span < self.LEADING_CONTEXT_MIN_SPAN:
            return False

        if next_row and next_row.cells and next_row.cells[0].text == first.text:
            return True

        return first.col_span >= max(cell.col_span for cell in row.cells[1:])

    def _calc_effective_cols(self, grid: list[list[str]]) -> int:
        """右端の空列を除いた有効列数を返す。

        方眼紙 Excel では見た目を整えるためだけの空列が多い。
        ここで対象列を絞り、後続の判定と表出力を簡潔にする。

        戻り値:
            データが存在する最大列番号 + 1（0ベース）
        """
        if not grid:
            return 0
        max_col = len(grid[0])
        for col_idx in range(max_col - 1, -1, -1):
            if any(grid[row_idx][col_idx] != "" for row_idx in range(len(grid))):
                return col_idx + 1
        return 0

    def _render_table_block(
        self,
        rows: list[RowProfile],
        file_path: str,
        sheet_name: str,
    ) -> tuple[str, ChunkMeta]:
        """連続した表行を Markdown テーブルへ変換する。

        実データがある列だけを取り出し、連続する重複行を除外してから
        先頭行をヘッダーとして扱う。

        引数:
            rows: テーブル構造と判定された連続するRowProfile群
            file_path: ソースファイル名（メタデータ用）
            sheet_name: シート名（メタデータ用）

        戻り値:
            markdown: Markdown表形式の文字列
            chunk: チャンクメタデータ
        """
        if not rows:
            return "", self._build_chunk_meta(file_path, ContentType.TABLE)

        active_columns = sorted({cell.start_col for row in rows for cell in row.cells})
        if not active_columns:
            return "", self._build_chunk_meta(file_path, ContentType.TABLE)

        compressed_rows = [
            [row.raw_row[col_idx] for col_idx in active_columns] for row in rows
        ]
        deduped_rows = self._dedupe_consecutive_rows(compressed_rows)
        if not deduped_rows:
            return "", self._build_chunk_meta(file_path, ContentType.TABLE)

        headers = deduped_rows[0]
        data_rows = deduped_rows[1:]
        markdown = MarkdownEscaper.make_table(headers, data_rows)

        col_start = get_column_letter(active_columns[0] + 1)
        col_end = get_column_letter(active_columns[-1] + 1)
        cell_range = (
            f"{col_start}{rows[0].row_index + 1}:{col_end}{rows[-1].row_index + 1}"
        )
        chunk = self._build_chunk_meta(
            file_path=file_path,
            content_type=ContentType.TABLE,
            sheet_name=sheet_name,
            cell_range=cell_range,
        )
        return markdown, chunk

    def _dedupe_consecutive_rows(self, rows: list[list[str]]) -> list[list[str]]:
        """連続して同じ内容の行を 1 行にまとめる。"""
        deduped: list[list[str]] = []
        previous_key: tuple[str, ...] | None = None
        for row in rows:
            normalized_key = tuple(MarkdownEscaper.escape_cell(cell) for cell in row)
            if all(cell == "" for cell in normalized_key):
                continue
            if normalized_key == previous_key:
                continue
            deduped.append(row)
            previous_key = normalized_key
        return deduped

    def _extract_comments(self, ws: Worksheet) -> str:
        """セルコメントを Markdown に変換する。

        openpyxl の制限:
        - 通常のコメント（セルメモ）は抽出可能

        戻り値:
            コメントを列挙したマークダウン（コメントがない場合は空文字列）
        """
        comments_md: list[str] = []
        for row in ws.iter_rows():
            for cell in row:
                if not cell.comment:
                    continue
                column_index = cell.column
                if column_index is None:
                    continue
                coord = f"{get_column_letter(column_index)}{cell.row}"
                comment_text = MarkdownEscaper.escape_cell(cell.comment.text)
                comments_md.append(f"- **[{coord}]** {comment_text}")

        if not comments_md:
            return ""
        return "### コメント・注記\n\n" + "\n".join(comments_md)
